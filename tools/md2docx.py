#!/usr/bin/env python3
"""md2docx.py - 简历 Markdown 转 Word（.docx）

支持的结构化子集（与 cv/main_master.md 对齐）：
  # / ## / ###   标题（1/2/3 级）
  - ...          无序列表项
  **bold**       行内加粗
  [text](url)    行内超链接（docx 中保留真实超链接）
  ---            分隔线（忽略）

用法: python3 tools/md2docx.py <input.md> <output.docx>
"""
import logging
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
log = logging.getLogger("md2docx")

FONT_CN = "Microsoft YaHei"
FONT_EN = "Microsoft YaHei"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)  # 深蓝，贴近 moderncv blue

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def set_run_font(run, size=None, bold=False, color=None):
    """同时设置西文与中文字体（python-docx 的 font.name 只覆盖西文）。"""
    run.font.name = FONT_EN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_hyperlink(paragraph, text, url, size):
    """在段落中追加真实超链接（ATS 与 Word 都可识别）。"""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = paragraph._p.makeelement(qn("w:hyperlink"), {qn("r:id"): r_id})
    paragraph._p.append(hyperlink)
    run = paragraph.add_run()
    run.text = text
    set_run_font(run, size=size, color=RGBColor(0x05, 0x63, 0xC1))
    run.underline = True
    hyperlink.append(run._element)


def add_inline(paragraph, text, size, base_bold=False):
    """解析 **bold** 与 [text](url)，按顺序写入段落。"""
    # 先按链接切分，链接内部不再解析加粗（简历文本中不存在该嵌套）
    pos = 0
    for m in LINK_RE.finditer(text):
        add_inline_plain(paragraph, text[pos:m.start()], size, base_bold)
        add_hyperlink(paragraph, m.group(1), m.group(2), size)
        pos = m.end()
    add_inline_plain(paragraph, text[pos:], size, base_bold)


def add_inline_plain(paragraph, text, size, base_bold):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            set_run_font(run, size=size, bold=base_bold)
        run = paragraph.add_run(m.group(1))
        set_run_font(run, size=size, bold=True)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, bold=base_bold)


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(46)
        section.left_margin = section.right_margin = Pt(50)

    stats = {"h1": 0, "h2": 0, "h3": 0, "bullet": 0, "para": 0, "link": 0}
    for raw in lines:
        line = raw.rstrip()
        if not line or line.strip() == "---":
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_plain(p, line[2:], size=22, base_bold=True)
            stats["h1"] += 1
        elif line.startswith("## "):
            p = doc.add_paragraph()
            add_inline_plain(p, line[3:], size=15, base_bold=True, )
            for run in p.runs:
                run.font.color.rgb = ACCENT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            stats["h2"] += 1
        elif line.startswith("### "):
            p = doc.add_paragraph()
            add_inline_plain(p, line[4:], size=12, base_bold=True)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            stats["h3"] += 1
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:], size=10.5)
            stats["bullet"] += 1
        else:
            p = doc.add_paragraph()
            add_inline(p, line, size=10.5)
            stats["para"] += 1
        stats["link"] += len(LINK_RE.findall(line))

    doc.save(docx_path)
    log.info("written %s: %s", docx_path, stats)


def main():
    if len(sys.argv) != 3:
        log.error("usage: md2docx.py <input.md> <output.docx>")
        sys.exit(2)
    convert(sys.argv[1], sys.argv[2])


if __name__ == "__main__":
    main()
